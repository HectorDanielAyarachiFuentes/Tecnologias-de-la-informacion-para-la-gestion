import openpyxl
from openpyxl.chart import BarChart, Reference
import urllib.request
import urllib.parse
import json
import csv
import os

# 1. Rutas
input_excel = "../Actividad - Propuesta/Final 2024/final2024.xlsx"
out_dir = "../Entregables/Final 2024"
out_excel = os.path.join(out_dir, "final2024.xlsx")
out_csv = os.path.join(out_dir, "final-TIC2024.csv")

if not os.path.exists(out_dir):
    os.makedirs(out_dir)

wb = openpyxl.load_workbook(input_excel)
ws = wb.active
ws.title = "Encuesta"

# Extraer todos los datos como lista de diccionarios para facilitar filtrado
headers = [cell.value for cell in ws[1]]
data = []
for row in ws.iter_rows(min_row=2, values_only=True):
    data.append(dict(zip(headers, row)))

ws["L1"] = "par-categoría"
ws["M1"] = "Mayúscula"
ws["N1"] = "Plus por cat"
headers.extend(["par-categoría", "Mayúscula", "Plus por cat"])

cat_counts = {}
reg_cat_counts = {}
reg_counts = {}

# Procesar filas y escribir fórmulas
for row_idx, row_data in enumerate(data, start=2):
    # Fórmulas de Excel
    ws.cell(row=row_idx, column=6, value=f"=2024-E{row_idx}") # Edad
    ws.cell(row=row_idx, column=12, value=f"=ISEVEN(H{row_idx})") # Par-categoría
    ws.cell(row=row_idx, column=13, value=f"=UPPER(C{row_idx})") # Mayúscula
    ws.cell(row=row_idx, column=14, value=f'=IF(H{row_idx}>=3, "Si corresponde", " ")') # Plus

    # Para los filtrados y tablas de Typst necesitamos calcular localmente
    cat = row_data.get('categoria')
    reg = row_data.get('regimen')
    edad = 2024 - row_data.get('año de nacimiento') if row_data.get('año de nacimiento') else 0
    
    if cat is not None:
        cat_counts[cat] = cat_counts.get(cat, 0) + 1
    if reg is not None and cat is not None:
        if reg not in reg_cat_counts:
            reg_cat_counts[reg] = {}
        reg_cat_counts[reg][cat] = reg_cat_counts[reg].get(cat, 0) + 1
    if reg is not None:
        reg_counts[reg] = reg_counts.get(reg, 0) + 1

# P7-a
p7a = wb.create_sheet("P7-a")
p7a.append(headers)
current_row = 2
for row_data in data:
    if row_data.get('region') == 'AN' and row_data.get('regimen') == 'Personal':
        new_row = [row_data.get(h, "") for h in headers]
        new_row[5] = f"=2024-E{current_row}"
        new_row[11] = f"=ISEVEN(H{current_row})"
        new_row[12] = f"=UPPER(C{current_row})"
        new_row[13] = f'=IF(H{current_row}>=3, "Si corresponde", " ")'
        p7a.append(new_row)
        current_row += 1

# P7-b
p7b = wb.create_sheet("P7-b")
p7b.append(headers)
current_row = 2
for row_data in data:
    edad_val = 2024 - row_data.get('año de nacimiento', 0) if row_data.get('año de nacimiento') else 0
    if edad_val < 30 and row_data.get('regimen') == 'Funcionarios':
        new_row = [row_data.get(h, "") for h in headers]
        new_row[5] = f"=2024-E{current_row}"
        new_row[11] = f"=ISEVEN(H{current_row})"
        new_row[12] = f"=UPPER(C{current_row})"
        new_row[13] = f'=IF(H{current_row}>=3, "Si corresponde", " ")'
        p7b.append(new_row)
        current_row += 1

# P7-c
p7c = wb.create_sheet("P7-c")
p7c.append(headers)
current_row = 2
for row_data in data:
    if row_data.get('regimen') == 'Funcionarios' and row_data.get('region') == 'CL':
        new_row = [row_data.get(h, "") for h in headers]
        new_row[5] = f"=2024-E{current_row}"
        new_row[11] = f"=ISEVEN(H{current_row})"
        new_row[12] = f"=UPPER(C{current_row})"
        new_row[13] = f'=IF(H{current_row}>=3, "Si corresponde", " ")'
        p7c.append(new_row)
        current_row += 1

# 9. Cantidad de empleados por régimen y gráfico de barras en Encuesta
start_res_row = ws.max_row + 2
ws.cell(row=start_res_row, column=1, value="Régimen")
ws.cell(row=start_res_row, column=2, value="Cantidad")

for i, (reg, count) in enumerate(reg_counts.items(), start=1):
    ws.cell(row=start_res_row+i, column=1, value=reg)
    ws.cell(row=start_res_row+i, column=2, value=count)

chart = BarChart()
chart.title = "Empleados por Régimen"
chart.x_axis.title = "Régimen"
chart.y_axis.title = "Cantidad"
data_ref = Reference(ws, min_col=2, min_row=start_res_row, max_row=start_res_row+len(reg_counts))
cats_ref = Reference(ws, min_col=1, min_row=start_res_row+1, max_row=start_res_row+len(reg_counts))
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats_ref)
ws.add_chart(chart, "A" + str(start_res_row + len(reg_counts) + 2))

wb.save(out_excel)

with open(out_csv, 'w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f)
    writer.writerow(headers)
    for row in ws.iter_rows(min_row=2, max_row=start_res_row-2, values_only=True):
        writer.writerow(row)

# Generar datos json para typst (en vez de gráficos directos, typst los dibujará o incrustaremos los generados)
# Usaremos quickchart para simplificar
cat_labels = list(cat_counts.keys())
cat_data = [cat_counts[k] for k in cat_labels]

pie_config = {
  "type": "pie",
  "data": {
    "labels": [f"Cat {l}" for l in cat_labels],
    "datasets": [{"data": cat_data}]
  }
}
pie_url = "https://quickchart.io/chart?c=" + urllib.parse.quote(json.dumps(pie_config))
urllib.request.urlretrieve(pie_url, os.path.join(out_dir, "pie_chart.png"))

regimens = list(reg_cat_counts.keys())
all_cats = list(set([c for reg in regimens for c in reg_cat_counts[reg].keys()]))
datasets = []
for c in all_cats:
    d = []
    for r in regimens:
        d.append(reg_cat_counts[r].get(c, 0))
    datasets.append({"label": f"Cat {c}", "data": d})

bar_config = {
  "type": "bar",
  "data": {
    "labels": regimens,
    "datasets": datasets
  }
}
bar_url = "https://quickchart.io/chart?c=" + urllib.parse.quote(json.dumps(bar_config))
urllib.request.urlretrieve(bar_url, os.path.join(out_dir, "bar_chart.png"))

print("Script completado. Archivos y gráficos generados.")
